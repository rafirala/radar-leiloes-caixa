"""Dá uma nota de oportunidade (0 a 85) a cada imóvel, com o motivo explicado, e gera os dados do mapa.

Nota (pesos simples e visíveis, para o usuário entender e discordar):
- até 40 pts: desconto sobre a avaliação da própria Caixa (60% ou mais = 40 pts). Compara o imóvel com ele mesmo.
- até 15 pts: preço do m² contra o preço médio do m² da cidade no FipeZAP (só casas e apartamentos em cidades cobertas).
  Contexto, não prova: a média da cidade inclui bairros nobres e prédios novos, e o imóvel da Caixa costuma ser popular.
O total vai até 85; "destaque" = 5% de maior nota; "conferir" = área acima de 300 m² ou preço acima de R$ 1 milhão.
- 15 pts: aceita financiamento.
- até 15 pts: modalidade (venda direta/online 15; licitação aberta 8; leilão SFI 0, que costuma começar pela avaliação).

Entrada: dados/imoveis_geo.csv (ou dados/imoveis.csv se ainda não localizou) e a planilha do FipeZAP.
Saída: docs/dados_mapa.json
Uso: python src/oportunidades.py
"""
import csv
import hashlib
import json
import re
import unicodedata
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parent.parent
FIPEZAP = RAIZ / "dados" / "bruto" / "fipezap_series.xlsx"  # https://downloads.fipe.org.br/indices/fipezap/fipezap-serieshistoricas.xlsx
MODALIDADE = {"Venda Direta Online": 15, "Venda Online": 15, "Licitação Aberta": 8}


def norm(s):
    return unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().upper().strip()


def mercado_fipezap():
    """{CIDADE: (mês, R$/m² médio de venda)} — último mês disponível de cada cidade."""
    wb = openpyxl.load_workbook(FIPEZAP, read_only=True, data_only=True)
    out = {}
    for aba in wb.sheetnames[3:]:
        linhas = [x for x in list(wb[aba].iter_rows(values_only=True))[4:]
                  if x and hasattr(x[1], "year") and isinstance(x[17], (int, float))]
        if linhas:
            out[norm(aba)] = (linhas[-1][1].strftime("%m/%Y"), round(linhas[-1][17]))
    return out


ABREV = {"AV": "Avenida", "AVENIDA": "Avenida", "R": "Rua", "RUA": "Rua", "AL": "Alameda", "TV": "Travessa", "TRAV": "Travessa",
         "EST": "Estrada", "ESTR": "Estrada", "ROD": "Rodovia", "PC": "Praça", "PCA": "Praça", "PRACA": "Praça", "LGO": "Largo",
         "VL": "Vila", "SERV": "Servidão", "ESTRADA": "Estrada", "RODOVIA": "Rodovia", "TRAVESSA": "Travessa", "ALAMEDA": "Alameda"}
MINUSC = {"DA", "DAS", "DE", "DO", "DOS", "E"}


def endereco_busca(e):
    """'AV ALFREDO RAVANELI, N. 3409, APTO 12' -> 'Avenida Alfredo Ravaneli, 3409' (para buscar no Google Maps).
    Tira 'N.', complementos (apto, bloco, lote...) e expande as abreviações da Caixa; 'SN' (sem número) fica sem número."""
    partes = [p.strip() for p in (e or "").upper().split(",")]
    if not partes or not partes[0]:
        return ""
    palavras = partes[0].replace(".", " ").split()
    if palavras and palavras[0] in ABREV:
        palavras[0] = ABREV[palavras[0]]
    rua = " ".join(w if w[0].isupper() and not w.isupper() else (w.lower() if w in MINUSC else w.capitalize()) for w in palavras)
    num = ""
    if len(partes) > 1:
        m = re.match(r"N\.?\s*(\d+)", partes[1])
        if m:
            num = m.group(1)
    return f"{rua}, {num}" if num else rua


def nota(x, merc):
    pts, motivos = 0, []
    d = float(x["desconto_pct"] or 0)
    p = min(d, 60) / 60 * 40
    pts += p
    motivos.append(f"{d:.0f}% abaixo da avaliação da Caixa (+{p:.0f})")
    ref = merc.get(norm(x["cidade"])) if x["tipo"] in ("Apartamento", "Casa") else None
    x["mercado_m2"] = ref[1] if ref else None
    if ref and x["preco_m2"]:
        r = float(x["preco_m2"]) / ref[1]
        p = max(0.0, min(1.0, (1 - r) / 0.6)) * 15
        pts += p
        motivos.append(f"m² a {r * 100:.0f}% da média da cidade no FipeZAP (+{p:.0f})")
    else:
        motivos.append("sem preço de mercado de referência para a cidade (+0)")
    if x["financiamento"] == "True":
        pts += 15
        motivos.append("aceita financiamento (+15)")
    m = MODALIDADE.get(x["modalidade"], 0)
    pts += m
    motivos.append(f"{x['modalidade']} (+{m})")
    return round(pts), motivos


def main():
    arq = RAIZ / "dados" / ("imoveis_geo.csv" if (RAIZ / "dados" / "imoveis_geo.csv").exists() else "imoveis.csv")
    imoveis = list(csv.DictReader(open(arq, encoding="utf-8")))
    merc = mercado_fipezap()
    saida = []
    for x in imoveis:
        n, motivos = nota(x, merc)
        num = lambda k: float(x[k]) if x.get(k) not in (None, "") else None
        lat, lon = num("lat"), num("lon")
        if lat is not None and x.get("precisao") in ("bairro", "cidade"):
            # vários imóveis no mesmo centro de bairro/cidade: espalha um pouco, sempre igual para o mesmo imóvel
            raio = 0.006 if x["precisao"] == "bairro" else 0.03
            h = int(hashlib.md5(x["id"].encode()).hexdigest(), 16)
            lat += ((h % 1000) / 1000 - 0.5) * raio
            lon += (((h // 1000) % 1000) / 1000 - 0.5) * raio
        saida.append({"id": x["id"], "uf": x["uf"], "cidade": x["cidade"].title(), "bairro": x["bairro"].title(), "endereco": x["endereco"].strip(), "endereco_busca": endereco_busca(x["endereco"]),
                      "tipo": x["tipo"], "preco": num("preco"), "avaliacao": num("avaliacao"), "desconto": num("desconto_pct"),
                      "area": num("area_privativa") or num("area_total"), "quartos": num("quartos"), "preco_m2": num("preco_m2"),
                      "mercado_m2": x["mercado_m2"], "financiamento": x["financiamento"] == "True", "modalidade": x["modalidade"],
                      "link": x["link"], "lat": lat, "lon": lon, "precisao": x.get("precisao", ""),
                      "nota": n, "motivos": motivos})
    saida.sort(key=lambda s: -s["nota"])
    corte = saida[max(0, len(saida) // 20 - 1)]["nota"] if saida else 0
    for s in saida:
        s["destaque"] = s["nota"] >= corte
        s["conferir"] = bool((s["area"] or 0) > 300 or (s["preco"] or 0) > 1_000_000)
    data_lista = imoveis[0].get("data_lista", "") if imoveis else ""
    (RAIZ / "docs").mkdir(exist_ok=True)
    (RAIZ / "docs" / "dados_mapa.json").write_text(json.dumps({"data_lista": data_lista, "imoveis": saida},
                                                              ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"{len(saida)} imóveis | localizados: {sum(1 for s in saida if s['lat'])} | destaques (nota ≥ {corte}): {sum(s['destaque'] for s in saida)} | conferir: {sum(s['conferir'] for s in saida)}")
    for s in saida[:8]:
        print(s["nota"], s["cidade"], s["bairro"], s["tipo"], s["preco"], "|", "; ".join(s["motivos"]))


if __name__ == "__main__":
    main()
